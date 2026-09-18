#!/usr/bin/env python3
"""Select representative book-marketing videos from XLSX or CSV data."""

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any


ALIASES = {
    "url": ["视频链接", "原视频链接", "链接", "视频地址", "抖音链接", "url", "URL"],
    "video_id": ["视频ID", "视频id", "作品ID", "作品id", "video_id"],
    "book": ["书名", "书籍", "图书", "商品", "商品名称", "推广图书", "带货商品"],
    "sales": ["销量", "视频销量", "预估销量", "商品销量", "成交件数", "销售量", "订单量", "总销量"],
    "gmv": ["销售额", "视频销售额", "成交金额", "预估销售额", "GMV", "gmv"],
    "likes": ["点赞", "点赞数", "获赞", "点赞量"],
    "comments": ["评论", "评论数", "评论量"],
    "favorites": ["收藏", "收藏数", "收藏量"],
    "shares": ["转发", "转发数", "分享", "分享数"],
    "title": ["标题", "视频标题", "文案", "作品标题"],
}


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def parse_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return 0.0 if isinstance(value, float) and math.isnan(value) else float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    if text in {"", "-", "--", "—", "无", "暂无", "N/A", "n/a"}:
        return 0.0
    multipliers = {"万": 10000.0, "w": 10000.0, "W": 10000.0, "千": 1000.0, "k": 1000.0, "K": 1000.0, "亿": 100000000.0}
    match = re.search(r"(-?\d+(?:\.\d+)?)\s*([万亿千wWkK]?)", text)
    if not match:
        return 0.0
    number, suffix = match.groups()
    return float(number) * multipliers.get(suffix, 1.0)


def resolve_columns(headers: list[str]) -> dict[str, str | None]:
    normalized = {normalize_header(h).lower(): h for h in headers}
    result: dict[str, str | None] = {}
    for key, names in ALIASES.items():
        result[key] = next(
            (normalized[normalize_header(name).lower()] for name in names if normalize_header(name).lower() in normalized),
            None,
        )
    return result


def read_xlsx(path: Path, sheet: str | None) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("读取 XLSX 需要 openpyxl。") from exc
    workbook = load_workbook(path, data_only=True, read_only=False)
    worksheets = [workbook[sheet]] if sheet else workbook.worksheets
    records: list[dict[str, Any]] = []
    all_headers: list[str] = []
    for worksheet in worksheets:
        rows = list(worksheet.iter_rows())
        if not rows:
            continue
        header_index = next(
            (index for index, row in enumerate(rows) if sum(cell.value not in (None, "") for cell in row) >= 2),
            None,
        )
        if header_index is None:
            continue
        headers = [
            str(cell.value).strip() if cell.value not in (None, "") else f"未命名列{index + 1}"
            for index, cell in enumerate(rows[header_index])
        ]
        for header in headers:
            if header not in all_headers:
                all_headers.append(header)
        for row in rows[header_index + 1 :]:
            values: list[Any] = []
            for cell in row[: len(headers)]:
                value = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else cell.value
                values.append(value)
            if not any(value not in (None, "") for value in values):
                continue
            record = dict(zip(headers, values))
            record["来源工作表"] = worksheet.title
            records.append(record)
    if "来源工作表" not in all_headers:
        all_headers.append("来源工作表")
    return records, all_headers


def read_csv(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                records = list(reader)
                return records, list(reader.fieldnames or [])
        except UnicodeDecodeError:
            continue
    raise SystemExit("无法识别 CSV 编码。")


def stable_key(record: dict[str, Any], columns: dict[str, str | None], index: int) -> str:
    for logical in ("video_id", "url"):
        column = columns.get(logical)
        if column and record.get(column):
            return str(record[column]).strip()
    title_column = columns.get("title")
    return f"title:{record.get(title_column, '')}:{index}"


def choose(
    records: list[dict[str, Any]],
    columns: dict[str, str | None],
    limit: int,
    per_book: int,
    max_per_book: int | None,
) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, record in enumerate(records):
        key = stable_key(record, columns, index)
        if key in seen:
            continue
        url_column = columns.get("url")
        if url_column and not str(record.get(url_column) or "").strip():
            continue
        seen.add(key)
        item = dict(record)
        metrics = {
            name: parse_number(item.get(columns[name])) if columns.get(name) else 0.0
            for name in ("sales", "gmv", "likes", "comments", "favorites", "shares")
        }
        item["_metrics"] = metrics
        item["_engagement"] = (
            metrics["likes"] + 3 * metrics["comments"] + 2 * metrics["favorites"] + 3 * metrics["shares"]
        )
        item["_rank"] = (metrics["sales"], metrics["gmv"], item["_engagement"])
        deduped.append(item)

    deduped.sort(key=lambda item: item["_rank"], reverse=True)
    if not deduped:
        return []

    book_column = columns.get("book")
    has_sales = any(item["_metrics"]["sales"] > 0 or item["_metrics"]["gmv"] > 0 for item in deduped)
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if book_column and has_sales:
        for item in deduped:
            book = str(item.get(book_column) or "").strip()
            if book:
                groups[book].append(item)

    selected: list[dict[str, Any]] = []
    selected_ids: set[int] = set()
    selected_by_book: dict[str, int] = defaultdict(int)
    if groups:
        ordered_books = sorted(groups, key=lambda name: groups[name][0]["_rank"], reverse=True)
        for level in range(per_book):
            for book in ordered_books:
                if level < len(groups[book]) and len(selected) < limit:
                    if max_per_book is not None and selected_by_book[book] >= max_per_book:
                        continue
                    item = groups[book][level]
                    selected.append(item)
                    selected_ids.add(id(item))
                    selected_by_book[book] += 1
    for item in deduped:
        if len(selected) >= limit:
            break
        if id(item) not in selected_ids:
            book = str(item.get(book_column) or "").strip() if book_column else ""
            if max_per_book is not None and book and selected_by_book[book] >= max_per_book:
                continue
            selected.append(item)
            selected_ids.add(id(item))
            if book:
                selected_by_book[book] += 1

    selected.sort(key=lambda item: item["_rank"], reverse=True)
    for index, item in enumerate(selected, start=1):
        book = str(item.get(book_column) or "").strip() if book_column else ""
        if groups and book:
            reason = f"{book}销售表现优先入选"
        elif has_sales:
            reason = "整体销售表现优先入选"
        else:
            reason = "账号互动表现优先入选"
        item["代表序号"] = index
        item["综合互动值"] = round(item["_engagement"], 2)
        item["筛选理由"] = reason
    return selected


def write_output(path: Path, selected: list[dict[str, Any]], headers: list[str]) -> None:
    output_headers = ["代表序号", "筛选理由", "综合互动值"] + [
        header for header in headers if header not in {"代表序号", "筛选理由", "综合互动值"}
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=output_headers, extrasaction="ignore")
        writer.writeheader()
        for item in selected:
            writer.writerow({key: value for key, value in item.items() if not key.startswith("_")})


def main() -> int:
    parser = argparse.ArgumentParser(description="从爆款视频表格选择代表视频。")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=30)
    parser.add_argument("--per-book", type=int, default=3, help="每本书优先纳入的数量，不是上限")
    parser.add_argument("--max-per-book", type=int, help="每本书严格最多入选数量")
    parser.add_argument("--sheet", help="仅处理指定工作表；默认处理全部工作表")
    args = parser.parse_args()

    if not args.input.is_file():
        parser.error(f"输入文件不存在：{args.input}")
    if args.limit < 1 or args.per_book < 1 or (args.max_per_book is not None and args.max_per_book < 1):
        parser.error("--limit、--per-book 和 --max-per-book 必须大于 0")
    if args.input.suffix.lower() in {".xlsx", ".xlsm"}:
        records, headers = read_xlsx(args.input, args.sheet)
    elif args.input.suffix.lower() == ".csv":
        records, headers = read_csv(args.input)
    else:
        parser.error("仅支持 XLSX、XLSM 或 CSV。")
    if not records:
        raise SystemExit("表格中没有可用数据。")

    columns = resolve_columns(headers)
    if not columns["url"] and not columns["video_id"]:
        raise SystemExit("未找到视频链接或视频 ID 列。")
    selected = choose(records, columns, args.limit, args.per_book, args.max_per_book)
    if not selected:
        raise SystemExit("去重后没有可用视频。")
    write_output(args.output, selected, headers)
    print(f"已选择 {len(selected)} 条代表视频：{args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
